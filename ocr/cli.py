import sys
import os
import json
import argparse
from pathlib import Path

# Fix Windows console UTF-8 encoding
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Ensure root path is in sys.path
root_dir = Path(__file__).resolve().parent.parent
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

from ocr.pipeline import DocumentOCRPipeline
from ocr.page_detector import OCRStrategy

def process_docx(file_path):
    import docx
    doc = docx.Document(file_path)
    blocks = []
    full_text_parts = []
    reading_order = 0
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        reading_order += 1
        blocks.append({
            "block_id": f"b_{reading_order}",
            "page_number": 1,
            "reading_order": reading_order,
            "block_type": "PARAGRAPH",
            "text": txt,
            "bounding_box": {"x0": 50, "y0": reading_order * 25, "x1": 550, "y1": (reading_order * 25) + 20},
            "confidence": 0.99,
            "font_size": 11,
            "font_family": "Calibri"
        })
        full_text_parts.append(txt)
        
    for t in doc.tables:
        for row in t.rows:
            cell_texts = [c.text.strip() for c in row.cells if c.text.strip()]
            if cell_texts:
                reading_order += 1
                row_str = " | ".join(cell_texts)
                blocks.append({
                    "block_id": f"b_{reading_order}",
                    "page_number": 1,
                    "reading_order": reading_order,
                    "block_type": "TABLE_ROW",
                    "text": row_str,
                    "bounding_box": {"x0": 50, "y0": reading_order * 25, "x1": 550, "y1": (reading_order * 25) + 20},
                    "confidence": 0.99,
                    "font_size": 10,
                    "font_family": "Calibri"
                })
                full_text_parts.append(row_str)
                
    full_text = "\n\n".join(full_text_parts)
    return {
        "document_name": os.path.basename(file_path),
        "total_pages": 1,
        "document_confidence": 0.99,
        "engine_used": "DOCX_NATIVE_EXTRACTOR",
        "pages": [{
            "page_number": 1,
            "provenance": "DOCX_TEXT",
            "confidence": 0.99,
            "blocks_count": len(blocks),
            "blocks": blocks,
            "text": full_text
        }],
        "full_text": full_text,
    }

def main():
    parser = argparse.ArgumentParser(description="Revana OCR & Document Layout Extraction CLI")
    parser.add_argument("input_file", help="Path to PDF, DOCX, or image file (PNG, JPG, TIFF, BMP, WEBP)")
    parser.add_argument("--strategy", choices=["AUTO", "FORCE", "SKIP"], default="AUTO", help="OCR strategy for PDFs")
    parser.add_argument("--output", help="Optional path to write JSON output to file")
    
    args = parser.parse_args()
    
    input_path = os.path.abspath(args.input_file)
    if not os.path.exists(input_path):
        error_res = {"error": f"File not found: {input_path}", "status": "FAIL"}
        sys.stdout.buffer.write((json.dumps(error_res) + "\n").encode('utf-8'))
        sys.exit(1)
        
    ext = Path(input_path).suffix.lower()
    image_exts = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp"}
    docx_exts = {".docx", ".doc"}
    
    pipeline = DocumentOCRPipeline()
    
    try:
        if ext == ".pdf":
            strategy = OCRStrategy[args.strategy]
            result = pipeline.process_pdf(input_path, strategy=strategy)
        elif ext in docx_exts:
            result = process_docx(input_path)
        elif ext in image_exts:
            result = pipeline.process_image(input_path)
        else:
            try:
                result = pipeline.process_pdf(input_path)
            except Exception:
                result = pipeline.process_image(input_path)
                
        result["status"] = "SUCCESS"
        result["file_path"] = input_path
        
        json_output = json.dumps(result, ensure_ascii=False, indent=2)
        
        if args.output:
            with open(args.output, "w", encoding="utf-8") as f:
                f.write(json_output)
        else:
            sys.stdout.buffer.write((json_output + "\n").encode('utf-8'))
            
    except Exception as e:
        error_res = {"error": str(e), "status": "FAIL", "file_path": input_path}
        sys.stdout.buffer.write((json.dumps(error_res) + "\n").encode('utf-8'))
        sys.exit(1)

if __name__ == "__main__":
    main()
